#!/usr/bin/env python3
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys

def convert_csv_to_xlsx(csv_file):
    # 读取CSV文件
    df = pd.read_csv(csv_file)
    
    # 创建Excel writer对象
    xlsx_file = csv_file.replace('.csv', '.xlsx')
    writer = pd.ExcelWriter(xlsx_file, engine='openpyxl')
    
    # 写入Excel
    df.to_excel(writer, sheet_name='Results', index=False)
    
    # 获取工作表
    worksheet = writer.sheets['Results']
    
    # 定义样式
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    centered = Alignment(horizontal='center', vertical='center')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 应用表头样式
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = centered
        cell.border = border
    
    # 设置列宽和应用样式到所有单元格
    for idx, col in enumerate(df.columns):
        column = get_column_letter(idx + 1)
        max_length = max(
            df[col].astype(str).apply(len).max(),
            len(str(col))
        )
        worksheet.column_dimensions[column].width = max_length + 2
        
        # 应用样式到数据单元格
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet[f"{column}{row}"]
            cell.alignment = centered
            cell.border = border
            
            # 为数值列设置数字格式
            if col in ['KIOPS', 'BW(MiB/s)']:
                cell.number_format = '#,##0.00'
            elif col == 'Cache Size':
                cell.number_format = '#,##0'
    
    # 冻结首行
    worksheet.freeze_panes = 'A2'
    
    # 保存文件
    writer.close()
    print(f"Created Excel file: {xlsx_file}")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python3 csv2xlsx.py <csv_file>")
        sys.exit(1)
    
    csv_file = sys.argv[1]
    convert_csv_to_xlsx(csv_file) 